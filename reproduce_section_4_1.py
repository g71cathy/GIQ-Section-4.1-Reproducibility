#!/usr/bin/env python3
"""Reproduce all P-M-V statistics and Figure 3 used in Section 4.1."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

try:
    import scipy
    from scipy.stats import chi2_contingency, fisher_exact
except ImportError:  # The documented fallback is mathematically equivalent for these 2 x 2 tests.
    scipy = None
    chi2_contingency = fisher_exact = None


FIGURE_PANELS = [
    ("P", "(a) Governance problems", ["P_POLI", "P_PROC", "P_INT", "P_ENF", "P_PREC", "P_VERT", "P_DQ", "P_MAINT"]),
    ("M", "(b) Governing mechanisms", ["M_WF", "M_DGI", "M_KR", "M_EW", "M_PLAT", "M_HUB", "M_DIALOG", "M_DOC"]),
    ("V", "(c) Public value claims", ["V_EFF", "V_EXP", "V_RISK", "V_QUAL", "V_EQU", "V_SEC", "V_CAP", "V_BIZ"]),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_table(workbook, sheet_name: str, header_label: str) -> list[dict[str, object]]:
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    header_index = next(index for index, row in enumerate(rows) if row and row[0] == header_label)
    headers = [str(value) if value is not None else "" for value in rows[header_index]]
    output = []
    for values in rows[header_index + 1:]:
        if not any(value is not None for value in values):
            continue
        output.append(dict(zip(headers, values)))
    return output


def read_workbook(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = read_table(workbook, "Public data", "Record ID")
    codebook = read_table(workbook, "Codebook", "Family")
    expected = read_table(workbook, "Results", "Family")
    workbook.close()
    required = ["Record ID", "Corpus", "Source ID", "Problem code", "Mechanism code",
                "Value code 1", "Value code 2", "Value code 3"]
    if not rows or any(column not in rows[0] for column in required):
        raise ValueError(f"Public data must contain: {required}")
    if {str(row["Corpus"]) for row in rows} != {"Policy", "Case"}:
        raise ValueError("Corpus must contain exactly Policy and Case.")
    return rows, codebook, expected


def present(row: dict[str, object], family: str, code: str) -> int:
    if family == "P":
        return int(row["Problem code"] == code)
    if family == "M":
        return int(row["Mechanism code"] == code)
    return int(code in {row["Value code 1"], row["Value code 2"], row["Value code 3"]})


def fisher_manual(table: list[list[int]]) -> float:
    a, b = table[0]
    c, d = table[1]
    row1, row2, col1 = a + b, c + d, a + c
    total = row1 + row2
    lower, upper = max(0, col1 - row2), min(row1, col1)
    denominator = math.comb(total, col1)
    probability = lambda x: math.comb(row1, x) * math.comb(row2, col1 - x) / denominator
    observed = probability(a)
    return min(1.0, sum(probability(x) for x in range(lower, upper + 1)
                        if probability(x) <= observed + 1e-15))


def narrative_unit_test(y_policy: np.ndarray, y_case: np.ndarray) -> tuple[str, float, float]:
    table = np.array([
        [int(y_policy.sum()), int(len(y_policy) - y_policy.sum())],
        [int(y_case.sum()), int(len(y_case) - y_case.sum())],
    ], dtype=int)
    expected = np.outer(table.sum(axis=1), table.sum(axis=0)) / table.sum()
    minimum = float(expected.min())
    if minimum < 5:
        p_value = float(fisher_exact(table, alternative="two-sided").pvalue) if fisher_exact else fisher_manual(table.tolist())
        return "Fisher's exact test (two-sided)", p_value, minimum
    if chi2_contingency:
        p_value = float(chi2_contingency(table, correction=True).pvalue)
    else:
        chi2 = float(np.sum((np.maximum(np.abs(table - expected) - 0.5, 0) ** 2) / expected))
        p_value = math.erfc(math.sqrt(chi2 / 2))
    return "Pearson chi-square with Yates correction", p_value, minimum


def clustered_logit(corpus: np.ndarray, clusters: list[str], y: np.ndarray) -> tuple[float, float, float, float]:
    """Binary logit coefficient with CR1 source-clustered sandwich covariance."""
    y0, y1 = y[corpus == 0], y[corpus == 1]
    p0, p1 = float(y0.mean()), float(y1.mean())
    if p0 in {0.0, 1.0} or p1 in {0.0, 1.0}:
        return math.nan, math.nan, math.nan, math.nan
    beta0 = math.log(p0 / (1 - p0))
    beta1 = math.log(p1 / (1 - p1)) - beta0
    fitted = np.where(corpus == 1, p1, p0)
    design = np.column_stack([np.ones(len(corpus)), corpus])
    bread = design.T @ ((fitted * (1 - fitted))[:, None] * design)
    scores: dict[str, np.ndarray] = {}
    for cluster, xi, yi, pi in zip(clusters, design, y, fitted):
        scores.setdefault(cluster, np.zeros(2))
        scores[cluster] += xi * (yi - pi)
    meat = sum(np.outer(score, score) for score in scores.values())
    inverse = np.linalg.inv(bread)
    covariance = inverse @ meat @ inverse
    g, n, k = len(scores), len(y), design.shape[1]
    covariance *= (g / (g - 1)) * ((n - 1) / (n - k))
    se = math.sqrt(float(covariance[1, 1]))
    p_value = math.erfc(abs(beta1 / se) / math.sqrt(2))
    return math.exp(beta1), math.exp(beta1 - 1.96 * se), math.exp(beta1 + 1.96 * se), p_value


def significance(p_value: float) -> str:
    if math.isnan(p_value):
        return "N/E"
    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"
    return "ns"


def compute(rows, codebook):
    policy = [row for row in rows if row["Corpus"] == "Policy"]
    case = [row for row in rows if row["Corpus"] == "Case"]
    ordered = policy + case
    corpus = np.array([0] * len(policy) + [1] * len(case), dtype=float)
    clusters = [str(row["Source ID"]) for row in ordered]
    results = []
    for definition in codebook:
        family, code = str(definition["Family"]), str(definition["Code"])
        y = np.array([present(row, family, code) for row in ordered], dtype=int)
        y0, y1 = y[:len(policy)], y[len(policy):]
        method, naive_p, minimum = narrative_unit_test(y0, y1)
        odds, lower, upper, clustered_p = clustered_logit(corpus, clusters, y)
        results.append({
            "Family": family,
            "Code": code,
            "Full term": str(definition["Full term"]),
            "Figure 3": str(definition["Figure 3"]),
            "Policy count": int(y0.sum()),
            "Policy N": len(y0),
            "Policy share": float(y0.mean()),
            "Case count": int(y1.sum()),
            "Case N": len(y1),
            "Case share": float(y1.mean()),
            "Minimum expected count": minimum,
            "Narrative-unit test": method,
            "Narrative-unit p": naive_p,
            "Case/Policy OR": odds,
            "95% CI lower": lower,
            "95% CI upper": upper,
            "Clustered p": clustered_p,
            "Significance": significance(clustered_p),
        })
    return results


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def choose_font(size: int, bold: bool = False):
    candidates = [
        Path("/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Times New Roman.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def plot_figure(results: list[dict[str, object]], output: Path) -> None:
    lookup = {str(row["Code"]): row for row in results}
    policy_color, case_color, accent = "#AABCC1", "#78B7B2", "#D28732"
    text_color, grid, panel = "#1F1F1F", "#E6E2DC", "#C9C6C0"
    width, height = 3900, 1380
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, code_font = choose_font(33, True), choose_font(25)
    value_font, value_bold, axis_font = choose_font(21), choose_font(21, True), choose_font(23)
    legend_font = choose_font(25)
    draw.rectangle((1270, 38, 1325, 78), fill=policy_color)
    draw.text((1340, 38), "Policy narratives", font=legend_font, fill=text_color)
    draw.rectangle((1690, 38, 1745, 78), fill=case_color)
    draw.text((1760, 38), "Case narratives", font=legend_font, fill=text_color)
    draw.text((2110, 38), "Orange asterisks: source-clustered p-value", font=legend_font, fill=accent)
    left_margin, panel_width, gap = 60, 1225, 52
    panel_top, panel_bottom = 130, 1265
    for panel_index, (_, panel_title, codes) in enumerate(FIGURE_PANELS):
        left = left_margin + panel_index * (panel_width + gap)
        right = left + panel_width
        draw.rounded_rectangle((left, panel_top, right, panel_bottom), radius=8, fill="#FBFBFA", outline=panel, width=2)
        title_box = draw.textbbox((0, 0), panel_title, font=title_font)
        draw.text(((left + right - (title_box[2] - title_box[0])) / 2, panel_top + 35), panel_title, font=title_font, fill=text_color)
        maximum = max(max(float(lookup[code]["Policy share"]), float(lookup[code]["Case share"])) for code in codes) * 100
        axis_max = max(20, int(math.ceil(maximum / 10) * 10))
        plot_left, plot_right = left + 225, right - 75
        plot_top, plot_bottom = panel_top + 135, panel_bottom - 105
        plot_width = plot_right - plot_left
        tick_step = 10 if axis_max <= 40 else 20
        for tick in range(0, axis_max + 1, tick_step):
            x = plot_left + plot_width * tick / axis_max
            draw.line((x, plot_top, x, plot_bottom), fill=grid, width=2)
            draw.text((x - 8, plot_bottom + 12), str(tick), font=value_font, fill=text_color)
        row_height = (plot_bottom - plot_top) / len(codes)
        for index, code in enumerate(codes):
            row = lookup[code]
            center_y = plot_top + row_height * (index + 0.5)
            box = draw.textbbox((0, 0), code, font=code_font)
            draw.text((plot_left - 20 - (box[2] - box[0]), center_y - 13), code, font=code_font, fill=text_color)
            for offset, field, color, add_stars in [(-19, "Policy share", policy_color, False), (20, "Case share", case_color, True)]:
                value = float(row[field]) * 100
                x2 = plot_left + plot_width * value / axis_max
                draw.rectangle((plot_left, center_y + offset - 12, x2, center_y + offset + 12), fill=color)
                star = str(row["Significance"]).replace("ns", "").replace("N/E", "") if add_stars else ""
                label = f"{value:.1f}{star}"
                draw.text((x2 + 8, center_y + offset - 12), label, font=value_bold if star else value_font,
                          fill=accent if star else text_color)
        axis_label = "Narrative units containing the element (%)"
        axis_box = draw.textbbox((0, 0), axis_label, font=axis_font)
        draw.text(((plot_left + plot_right - (axis_box[2] - axis_box[0])) / 2, panel_bottom - 60), axis_label,
                  font=axis_font, fill=text_color)
    image.save(output, dpi=(300, 300), optimize=True)


def verify(results, expected) -> dict[str, object]:
    expected_lookup = {(str(row["Family"]), str(row["Code"])): row for row in expected}
    numeric_fields = ["Policy count", "Policy N", "Policy share", "Case count", "Case N", "Case share",
                      "Narrative-unit p", "Case/Policy OR", "95% CI lower", "95% CI upper", "Clustered p"]
    maximum_difference = 0.0
    mismatches = []
    for result in results:
        key = (str(result["Family"]), str(result["Code"]))
        expected_row = expected_lookup.get(key)
        if expected_row is None:
            mismatches.append(f"Missing expected result: {key}")
            continue
        for field in numeric_fields:
            actual = math.nan if result[field] is None else float(result[field])
            target = math.nan if expected_row[field] is None else float(expected_row[field])
            if math.isnan(actual) and math.isnan(target):
                continue
            difference = abs(actual - target)
            maximum_difference = max(maximum_difference, difference)
            if difference > 1e-10:
                mismatches.append(f"{key} {field}: {actual} != {target}")
    return {"status": "PASS" if not mismatches else "FAIL", "maximum_absolute_difference": maximum_difference,
            "mismatches": mismatches}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Public_PMV_Element_Data.xlsx")
    parser.add_argument("--output-dir", type=Path, default=Path("reproduced_section_4_1"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    rows, codebook, expected = read_workbook(args.workbook)
    results = compute(rows, codebook)
    for family, filename in [("P", "Table_A1_Problems.csv"), ("M", "Table_A2_Mechanisms.csv"), ("V", "Table_A3_Values.csv")]:
        write_csv(args.output_dir / filename, [row for row in results if row["Family"] == family])
    write_csv(args.output_dir / "All_PMV_results.csv", results)
    plot_figure(results, args.output_dir / "Figure_3_reproduced.png")
    verification = verify(results, expected)
    (args.output_dir / "verification.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
    metadata = {
        "input_file": args.workbook.name,
        "input_sha256": sha256(args.workbook),
        "python": platform.python_version(),
        "numpy": np.__version__,
        "scipy": scipy.__version__ if scipy else "not installed; equivalent 2 x 2 fallback used",
        "openpyxl": openpyxl.__version__,
        "policy_narrative_units": sum(row["Corpus"] == "Policy" for row in rows),
        "case_narrative_units": sum(row["Corpus"] == "Case" for row in rows),
        "value_rule": "A value is present if it occurs in Value code 1, 2, or 3; duplicates count once per narrative unit.",
        "test_rule": "Two-sided Fisher exact test when any expected cell is below 5; otherwise Pearson chi-square with Yates correction.",
        "primary_inference": "Binary logit Case versus Policy with CR1 standard errors clustered by anonymous source ID.",
        "verification": verification["status"],
    }
    (args.output_dir / "run_metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(json.dumps(metadata, indent=2))
    if verification["status"] != "PASS":
        raise RuntimeError("Recomputed results do not match the workbook Results sheet.")


if __name__ == "__main__":
    main()
